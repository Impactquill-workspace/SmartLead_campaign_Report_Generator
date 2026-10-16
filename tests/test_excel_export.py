import os
import openpyxl
from app.utils.export_utils import generate_excel_report


def test_excel_export_generation(tmp_path):
    report_data = {
        "period_start": "2026-09-21",
        "period_end": "2026-09-27",
        "date_range_str": "2026-09-21 to 2026-09-27",
        "status_filter": "ACTIVE",
        "total_campaigns": 3,
        "total_sent": 850,
        "total_bounced": 12,
        "overall_bounce_rate": 1.41,
        "campaigns": [
            {
                "date_range": "2026-09-21 to 2026-09-27",
                "campaign_id": "camp_1",
                "campaign_name": "Client A Campaign 1",
                "total_sent": 300,
                "total_bounced": 5,
                "total_invalid": 1,
                "total_blocked": 2,
                "total_replies": 9,
                "positive_replies": 3,
                "bounce_percentage": 1.11,
                "total_leads": 450,
                "mimecast_bounces": 2,
            },
            {
                "date_range": "2026-09-21 to 2026-09-27",
                "campaign_id": "camp_2",
                "campaign_name": "Client A Campaign 2",
                "total_sent": 200,
                "total_bounced": 3,
                "total_invalid": 2,
                "total_blocked": 1,
                "total_replies": 4,
                "positive_replies": 1,
                "bounce_percentage": 0.94,
                "total_leads": 320,
                "mimecast_bounces": 1,
            },
            {
                "date_range": "2026-09-21 to 2026-09-27",
                "campaign_id": "camp_3",
                "campaign_name": "Client B Campaign 3",
                "total_sent": 350,
                "total_bounced": 4,
                "total_invalid": 0,
                "total_blocked": 0,
                "total_replies": 5,
                "positive_replies": 2,
                "bounce_percentage": 0.80,
                "total_leads": 500,
                "mimecast_bounces": 0,
            },
        ],
    }

    file_path = generate_excel_report(report_data, output_dir=str(tmp_path))
    assert os.path.exists(file_path)

    wb = openpyxl.load_workbook(file_path)
    sheet_names = wb.sheetnames
    assert "Campaign Delivery Report" in sheet_names
    assert "Report Summary" in sheet_names

    # Check Sheet 1 (Campaign Delivery Report) — 13 columns
    ws1 = wb["Campaign Delivery Report"]
    headers = [cell.value for cell in ws1[1]]
    assert headers == [
        "Date Range",
        "Campaign Name",
        "Total Leads",
        "Total sent",
        "Total bounce no",
        "Bounce percentage",
        "Blocked Emails",
        "Invalid Emails",
        "Reply Count",
        "Reply percentage",
        "Positive Replies",
        "Positive Reply percentage",
        "Mimecast Bounces",
    ]

    # Row 2 should be first campaign
    assert ws1.cell(row=2, column=1).value == "2026-09-21 to 2026-09-27"
    assert ws1.cell(row=2, column=2).value == "Client A Campaign 1"
    assert ws1.cell(row=2, column=3).value == 450  # Total Leads
    assert ws1.cell(row=2, column=4).value == 300  # Total sent
    assert ws1.cell(row=2, column=5).value == 5    # Total bounce no
    assert ws1.cell(row=2, column=7).value == 2    # Blocked Emails
    assert ws1.cell(row=2, column=8).value == 1    # Invalid Emails
    assert ws1.cell(row=2, column=9).value == 9    # Reply Count
    assert round(ws1.cell(row=2, column=10).value, 4) == round(9 / (300 - 5), 4)  # Reply percentage (9 / 295)
    assert ws1.cell(row=2, column=11).value == 3   # Positive Replies
    assert round(ws1.cell(row=2, column=12).value, 4) == round(3 / 9, 4)          # Positive Reply % (3 / 9)
    assert ws1.cell(row=2, column=13).value == 2   # Mimecast Bounces

    # Check Total row at row 5 (header + 3 data rows + 1 total row)
    assert ws1.cell(row=5, column=1).value == "Total"
    assert ws1.cell(row=5, column=3).value == 1270  # Total Leads sum
    assert ws1.cell(row=5, column=4).value == 850   # Total sent sum
    assert ws1.cell(row=5, column=5).value == 12    # Total bounce sum
    assert ws1.cell(row=5, column=7).value == 3     # Total blocked sum
    assert ws1.cell(row=5, column=8).value == 3     # Total invalid sum
    assert ws1.cell(row=5, column=9).value == 18    # Total replies sum
    assert round(ws1.cell(row=5, column=10).value, 4) == round(18 / (850 - 12), 4)
    assert ws1.cell(row=5, column=11).value == 6    # Total positive sum
    assert round(ws1.cell(row=5, column=12).value, 4) == round(6 / 18, 4)
    assert ws1.cell(row=5, column=13).value == 3    # Total mimecast sum

    # Check Sheet 2 (Report Summary)
    ws2 = wb["Report Summary"]
    assert ws2.max_row >= 12


